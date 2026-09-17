# Bot_Quantitativo.py
# Extensão do Bot Controle: cria/atualiza a aba "Quantitativo" na planilha de controle
# Regras (novo escopo):
# - Lê itens.txt (BOM Codes válidos)
# - Abre a cotação na aba "Lista Equipamentos WDM"
# - Para cada linha: compara BOM Code (coluna C) com itens.txt
# - Se bater, pega a quantidade total na coluna N (mesma linha)
# - Atualiza a aba "Quantitativo" (A: Gestor | B: ETP | C: BOM Code | D: Quantidade)
# - Sem gráfico

from __future__ import annotations

import os
import unicodedata
from typing import Optional, Set, Dict

from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Border, Side, Font


# =========================
# Utils
# =========================

def remover_acentos(texto: str) -> str:
    return unicodedata.normalize("NFKD", str(texto)).encode("ASCII", "ignore").decode("utf-8")


def normalizar_etp(etp: Optional[str]) -> str:
    """Normaliza ETP para 1234-5678 quando possível."""
    if not etp:
        return ""
    s = "".join(ch for ch in str(etp) if ch.isdigit())
    if len(s) == 8:
        return f"{s[:4]}-{s[4:]}"
    return s or str(etp).strip()


def _to_number(x) -> float:
    """Converte valores para float de forma tolerante (inteiro, float, string)."""
    if x is None:
        return 0.0
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip()
    if not s:
        return 0.0
    # tolerância BR: 1.234,56
    s = s.replace(" ", "").replace("R$", "")
    s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except Exception:
        return 0.0


# =========================
# itens.txt
# =========================

def carregar_itens_txt(base_dir: Optional[str] = None, nome_arquivo: str = "itens.txt") -> Set[str]:
    """
    Carrega BOM codes do itens.txt.
    Aceita os seguintes caminhos (em ordem):
      1) base_dir/nome_arquivo (se base_dir fornecido)
      2) pasta do próprio Bot_Quantitativo.py
      3) diretório atual (cwd)
    Ignora linhas vazias e comentários (#).
    """
    candidatos = []

    if base_dir:
        candidatos.append(os.path.join(base_dir, nome_arquivo))

    # pasta do arquivo atual
    candidatos.append(os.path.join(os.path.dirname(__file__), nome_arquivo))

    # cwd
    candidatos.append(os.path.join(os.getcwd(), nome_arquivo))

    path_ok = None
    for p in candidatos:
        if os.path.exists(p):
            path_ok = p
            break

    if not path_ok:
        raise FileNotFoundError(f"Não achei '{nome_arquivo}' nos caminhos: {candidatos}")

    itens: Set[str] = set()
    with open(path_ok, "r", encoding="utf-8", errors="ignore") as f:
        for line in f:
            raw = line.strip()
            if not raw or raw.startswith("#"):
                continue
            itens.add(raw)

    if not itens:
        raise ValueError(f"'{nome_arquivo}' foi encontrado, mas está vazio/sem itens válidos: {path_ok}")

    print(f"[Quant] itens.txt OK: {len(itens)} itens | path={path_ok}")
    return itens


# =========================
# Estilo/Format
# =========================

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

HEADER_FILL = PatternFill(start_color="B4C6E7", end_color="B4C6E7", fill_type="solid")
HEADER_FONT = Font(bold=True)


def _formatar_cabecalho(ws) -> None:
    headers = ["Gestor", "ETP", "BOM Code", "Quantidade"]
    for col, texto in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=texto)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER


def _aplicar_borda_tabela(ws, max_row: int, max_col: int) -> None:
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            ws.cell(row=r, column=c).border = THIN_BORDER


# =========================
# Extração da cotação
# =========================

def extrair_quantitativo_do_anexo(caminho_anexo_cotacao: str, itens_validos: Set[str]) -> Dict[str, float]:
    """
    Lê 'Lista Equipamentos WDM':
      - BOM Code: coluna C (3)
      - Total: coluna N (14)

    Retorna dict {BOM_code: quantidade_total}
    - Soma se o mesmo code aparecer mais de uma vez
    - Ignora qty <= 0
    - Ignora qualquer code fora de itens_validos
    """
    wb = None
    try:
        wb = load_workbook(caminho_anexo_cotacao, data_only=True)
    except Exception as e:
        print(f"[Quant][ERRO] Não foi possível abrir cotação: {caminho_anexo_cotacao} | {e}")
        return {}

    try:
        if "Lista Equipamentos WDM" not in wb.sheetnames:
            print(f"[Quant][ERRO] Aba 'Lista Equipamentos WDM' não encontrada em {os.path.basename(caminho_anexo_cotacao)}")
            return {}

        ws = wb["Lista Equipamentos WDM"]

        BOM_COL = 3   # C
        QTD_COL = 14  # N

        totais: Dict[str, float] = {}

        for row in range(1, ws.max_row + 1):
            code = ws.cell(row=row, column=BOM_COL).value
            if not code:
                continue

            code = str(code).strip()
            if code not in itens_validos:
                continue

            qtd = _to_number(ws.cell(row=row, column=QTD_COL).value)
            if qtd <= 0:
                continue

            totais[code] = totais.get(code, 0.0) + qtd

        return totais

    finally:
        try:
            wb.close()
        except Exception:
            pass


# =========================
# Atualização da planilha controle
# =========================

def atualizar_aba_quantitativo(
    caminho_planilha_controle: str,
    etp_display: Optional[str],
    etp_norm: Optional[str],
    gestor: Optional[str],
    quantitativo: Dict[str, float],
) -> None:
    """
    Atualiza a aba 'Quantitativo' na planilha de controle.

    Estratégia anti-duplicação:
      - Remove linhas existentes da MESMA ETP (comparando etp_norm)
      - Insere novamente os itens atuais (code->qtd)
    """
    if not quantitativo:
        print("[Quant] quantitativo vazio, nada a atualizar.")
        return

    etp_norm_local = normalizar_etp(etp_norm or etp_display)
    if not etp_norm_local:
        print("[Quant][ERRO] ETP não informada (etp_display/etp_norm vazios).")
        return

    wb = load_workbook(caminho_planilha_controle)
    try:
        if "Quantitativo" in wb.sheetnames:
            ws = wb["Quantitativo"]
            # garante cabeçalho
            if ws.max_row < 1 or ws["A1"].value is None:
                _formatar_cabecalho(ws)
        else:
            ws = wb.create_sheet("Quantitativo")
            _formatar_cabecalho(ws)

        # 1) Remove linhas existentes da mesma ETP
        # (iterar de baixo pra cima pra não bagunçar índices)
        removed = 0
        for r in range(ws.max_row, 1, -1):
            b = ws.cell(row=r, column=2).value  # ETP col B
            if normalizar_etp(b) == etp_norm_local:
                ws.delete_rows(r, 1)
                removed += 1

        if removed:
            print(f"[Quant] Removidas {removed} linhas antigas da ETP {etp_norm_local} (evita duplicação)")

        # 2) Insere itens novos
        next_row = ws.max_row + 1
        etp_to_write = etp_display or etp_norm_local
        gestor_to_write = gestor or ""

        inserted = 0
        for code in sorted(quantitativo.keys()):
            qtd = quantitativo[code]
            if qtd <= 0:
                continue

            ws.cell(row=next_row, column=1, value=gestor_to_write)
            ws.cell(row=next_row, column=2, value=etp_to_write)
            ws.cell(row=next_row, column=3, value=code)
            # se for inteiro "perfeito", grava como int (fica mais bonito)
            if abs(qtd - int(qtd)) < 1e-9:
                ws.cell(row=next_row, column=4, value=int(qtd))
            else:
                ws.cell(row=next_row, column=4, value=float(qtd))

            next_row += 1
            inserted += 1

        # 3) Bordas na área usada (A:D)
        max_row = ws.max_row
        _aplicar_borda_tabela(ws, max_row=max_row, max_col=4)

        # 4) Cabeçalho com estilo (garante)
        _formatar_cabecalho(ws)

        wb.save(caminho_planilha_controle)
        print(f"[Quant] OK | ETP={etp_norm_local} | Inseridos={inserted} | Planilha atualizada ✅")

    finally:
        try:
            wb.close()
        except Exception:
            pass


# =========================
# Função principal chamada pelo Bot Controle
# =========================

def processar_quantitativo(
    caminho_planilha_controle: str,
    etp_display: Optional[str],
    etp_norm: Optional[str],
    gestor: Optional[str],
    caminho_anexo_cotacao: str,
    base_dir_itens: Optional[str] = None,
) -> None:
    """
    Função PRINCIPAL (chamada pelo Bot Controle):

    - Carrega itens.txt
    - Lê a cotação (Lista Equipamentos WDM)
    - Compara col C vs itens.txt
    - Pega quantidade total na col N
    - Atualiza a aba Quantitativo
    """
    etp_label = etp_display or etp_norm or "(sem ETP)"
    print(f"[Quant] INÍCIO | ETP={etp_label} | Anexo={os.path.basename(caminho_anexo_cotacao or '')}")

    # 1) itens.txt
    try:
        itens_validos = carregar_itens_txt(base_dir=base_dir_itens, nome_arquivo="itens.txt")
    except Exception as e:
        print(f"[Quant][ERRO] Erro ao carregar itens.txt: {e}")
        return

    # 2) valida anexo
    if not caminho_anexo_cotacao or not os.path.exists(caminho_anexo_cotacao):
        print(f"[Quant][ERRO] Anexo de cotação não existe: {caminho_anexo_cotacao}")
        return

    # 3) extrai
    quantitativo = extrair_quantitativo_do_anexo(caminho_anexo_cotacao, itens_validos)
    if not quantitativo:
        print(f"[Quant] Nenhum item do itens.txt encontrado na cotação: {os.path.basename(caminho_anexo_cotacao)}")
        return

    print(f"[Quant] Encontrados {len(quantitativo)} BOM Codes na cotação (col C) com qtd em N")

    # 4) atualiza planilha controle
    try:
        atualizar_aba_quantitativo(
            caminho_planilha_controle=caminho_planilha_controle,
            etp_display=etp_display,
            etp_norm=etp_norm,
            gestor=gestor,
            quantitativo=quantitativo,
        )
    except Exception as e:
        print(f"[Quant][ERRO] Falha ao atualizar a aba Quantitativo: {e}")
        return

    print(f"[Quant] FIM | ETP={etp_label} ✅")


# =========================
# Execução local (opcional)
# =========================
if __name__ == "__main__":
    # Exemplo rápido de teste manual (ajuste caminhos):
    # processar_quantitativo(
    #     caminho_planilha_controle=r"C:\...\Controle.xlsx",
    #     etp_display="ETP 1234-5678 V1",
    #     etp_norm="1234-5678",
    #     gestor="Fulano",
    #     caminho_anexo_cotacao=r"C:\...\Cotacao_ETP_1234-5678.xlsx",
    #     base_dir_itens=os.path.dirname(__file__),
    # )
    pass
