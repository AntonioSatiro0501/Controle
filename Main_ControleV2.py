import os
import re
import time
import unicodedata
import win32com.client
import tkinter as tk
from tkinter import ttk
from tkinter import messagebox
from openpyxl import load_workbook
import threading
from datetime import datetime, timedelta
from Bot_Quantitativo import processar_quantitativo
import pythoncom
from excel_xlsxwriter import ler_planilhas, escrever_planilhas

# ==================== CONFIGURAÇÕES ====================

user_data = {}
path = os.path.dirname(__file__)
userDataPath = os.path.join(path, 'userData.txt')

with open(userDataPath, 'r', encoding='utf-8', errors='ignore') as file:
    lines = file.readlines()
    for line in lines:
        if ' = ' in line:
            key, value = line.strip().split(' = ', 1)
            user_data[key] = value

CAMINHO_PLANILHA_CONTROLE = user_data["Controle"]
PASTA_TEMP_ANEXOS = user_data["Anexos_Temp_ETPs"]
path_icone = user_data.get("Icone", "")

# Flags globais controladas pela UI (checkboxes)
WRITE_DISPLAY_TO_SHEET = True   # True: grava coluna ETP como na cotação; False: grava normalizada
USE_RESTRICT = True             # True: usa Items.Restrict (rápido); False: varre toda a caixa localmente

# Filtro diário (quando None, usa mensal)
DIA_FILTRADO = None

def remover_acentos(texto):
    return unicodedata.normalize('NFKD', texto).encode('ASCII', 'ignore').decode('utf-8')

# ---------------- RESTO DA CONFIG ---------------- #

_PADROES = {
    "segue a cotacao solicitada": "Enviado",
    "favor aprovar para envio da t2": "Enviado",
    "favor encaminhar a t2": "Pending T2",
    "favor enviar a t2": "Pending T2",
    "envia a t2 para a emissao dos pedidos": "Pending T2",
    "favor seguir com as t2": "Pending T2",
    "segue t2, conforme solicitado.": "T2 Enviado",
    "segue t2's, conforme solicitado.": "T2 Enviado",
    "segue arquivo corrigido.": "T2 Enviado",
    "por gentileza, considerar estas t2's para a etp.": "T2 Enviado",
    "por gentila, considerar estas t2's para a etp.": "T2 Enviado",
    "segue anexo correto para t2": "T2 Enviado",
    "segue t2": "T2 Enviado",
    "a cotacao esta de acordo com a ga, favor seguir com a t2.": "Pending T2",
    "a cotacao esta de acordo com a ga favor seguir com a t2": "Pending T2",
}
PADROES_CORPO = { remover_acentos(k.lower()): v for k, v in _PADROES.items() }

# Variáveis globais de contagem
TOTAL_EMAILS_ENCONTRADOS = 0
EMAILS_FILTRADOS = 0
PROCESSADOS = 0
VALIDADOS = 0
ADICIONADAS = 0
ATUALIZADAS = 0

# Listas detalhadas: (etp_display_exato, linha)
ADICIONADAS_LIST = []
ATUALIZADAS_LIST = []

MES_INICIO_CHAVES = 6       # mês em que foi aplicado o @Controle
ANO_INICIO_CHAVES = 2025    # ano em que foi aplicado o @Controle

# Constantes de coluna na planilha "Controle"
COL_GESTOR  = 1
COL_ETP     = 4
COL_STATUS  = 2
COL_ESTADOS = 7
COL_QNT     = 15 #A SER RETIRADO
COL_PRECO   = 16
COL_VAL_R   = 18
COL_VAL_S   = 19

# Variáveis para formatação de células
COL_LAST_BORDERED = 20

#======== Regras de status ========#
STATUS_ORDER = [
    "pending",      # [P]
    "enviado",      # [EV]
    "pending t2",   # [PT2]
    "t2 enviado",   # [ET2] (vem do anexo T2_ETP)
    "po received",  # Vem do BOT PREPO
]
STATUS_RANK = {name: i for i, name in enumerate(STATUS_ORDER, start=1)}

#======== Formatação dos dados quando adicionados no Excel ========#
BRL_ACCOUNTING = '[$R$-416] #,##0.00'

def _norm_status(s: str | None) -> str:
    if not s:
        return ""
    return remover_acentos(str(s)).strip().lower()

def _merge_status(existing: str | None, incoming: str | None) -> tuple[str, bool, str]:
    e = _norm_status(existing)
    n = _norm_status(incoming)

    if e == "po received":
        return existing, False, "PO Received é imutável"

    if not n:
        return existing, False, "Sem novo status"

    er = STATUS_RANK.get(e, 0)
    nr = STATUS_RANK.get(n, 0)

    if nr == 0:
        return existing, False, f"Novo status '{incoming}' desconhecido"

    if er >= nr:
        return existing, False, f"Não retrocede ({existing} ≥ {incoming})"

    return incoming, True, f"Progrediu: {existing or '(vazio)'} → {incoming}"

def _to_float(x):
    if x is None: return None
    if isinstance(x, (int, float)): return float(x)
    s = str(x).strip().replace('R$', '').replace(' ', '')
    s = s.replace('.', '').replace(',', '.')
    try:
        return float(s)
    except:
        return None

def status_por_chave_de_controle(corpo):
    if not corpo:
        return None
    corpo_norm = remover_acentos(corpo.lower()).replace(" ", "")
    if "@controle" not in corpo_norm:
        return None
    if "[pt2]" in corpo_norm: return "Pending T2"
    if "[ev]" in corpo_norm:  return "Enviado"
    if "[p]"  in corpo_norm:  return "Pending"
    return None

def normalizar_etp(etp):
    if not etp:
        return ""
    etp = re.sub(r'\D', '', str(etp))
    return f"{etp[:4]}-{etp[4:]}" if len(etp) == 8 else etp.strip()

def extrair_etp_do_assunto(assunto):
    if not assunto:
        return None
    a = remover_acentos(assunto).lower()
    m = re.search(r"\betp[\s\-_]*((\d{4})[\s\-_]?(\d{4}))\b", a)
    if m:
        return f"{m.group(2)}-{m.group(3)}"
    m2 = re.search(r"\betp[\s\-_]*(\d{8})\b", a)
    if m2:
        d = m2.group(1)
        return f"{d[:4]}-{d[4:]}"
    return None

def detectar_anexo_t2(paths):
    pad = re.compile(r"^t2[_\s-]*etp[_\s-]*(\d{4})[-_\s]?(\d{4}).*\.(xlsx|xlsm)$", re.IGNORECASE)
    for p in paths:
        base = remover_acentos(os.path.basename(p).lower())
        m = pad.match(base)
        if m:
            return True, f"{m.group(1)}-{m.group(2)}", p
    return False, None, None

def validar_email(assunto, anexos, corpo):
    if assunto and "ETP" in assunto.upper():
        return True
    if "@controle" in remover_acentos((corpo or "")).lower():
        return True
    for an in anexos:
        nome = remover_acentos(os.path.basename(an).lower())
        if nome.endswith(('.xlsx', '.xlsm')) and ('cotacao' in nome or 'etp' in nome or nome.startswith('t2_etp')):
            return True
    return False

def testar_conexao_outlook():
    try:
        pythoncom.CoInitialize()
        outlook = win32com.client.Dispatch('outlook.application').GetNamespace("MAPI")
        inbox = outlook.GetDefaultFolder(6)  # 6 = Inbox
        print(" Conexão com Outlook bem-sucedida.")
        return inbox
    except Exception as e:
        print(f" Erro ao conectar no Outlook: {e}")
        return None

def salvar_anexos(mensagem):
    anexos_salvos = []
    os.makedirs(PASTA_TEMP_ANEXOS, exist_ok=True)
    for anexo in mensagem.Attachments:
        nome = anexo.FileName
        nome_norm = remover_acentos(nome.lower())
        if nome_norm.endswith(('.xlsx', '.xlsm')) and ("cotacao" in nome_norm or "etp" in nome_norm or nome_norm.startswith("t2_etp")):
            caminho = os.path.join(PASTA_TEMP_ANEXOS, nome)
            anexo.SaveAsFile(caminho)
            anexos_salvos.append(caminho)
            print(f" Anexo salvo: {caminho}")
        else:
            print(f" Anexo ignorado: {nome}")
    return anexos_salvos

def identificar_status_etp(assunto, anexos, corpo, data_recebida, tem_t2=False):
    if tem_t2:
        return "T2 Enviado"

    corpo_norm = remover_acentos(corpo.lower())
    usa_chave = (
        data_recebida.year > ANO_INICIO_CHAVES or
        (data_recebida.year == ANO_INICIO_CHAVES and data_recebida.month >= MES_INICIO_CHAVES)
    )
    if usa_chave:
        s = status_por_chave_de_controle(corpo_norm)
        if s:
            print(f"Status via chave personalizada: {s}")
            return s
    else:
        for frase, status in PADROES_CORPO.items():
            if status == "T2 Enviado":
                continue 
            if frase in corpo_norm:
                print(f"Frase identificada: '{frase}' → {status}")
                return status
    print("Status não identificado.")
    return None

def extrair_dados_anexo(path_arquivo, etp_esperada_norm):
    try:
        wb = load_workbook(path_arquivo, data_only=True)

        aba_resumo = wb["Resumo"]
        etp_display = aba_resumo["B5"].value.strip() if aba_resumo["B5"].value else None
        etp_norm = normalizar_etp(etp_display) if etp_display else etp_esperada_norm

        gestor = aba_resumo["B6"].value.strip() if aba_resumo["B6"].value else None
        if isinstance(gestor, str) and gestor.lower().startswith("gestor tlf"):
            gestor = gestor.replace("Gestor TLF:", "").strip()

        preco = _to_float(aba_resumo["E7"].value)
        valor_r = sum((_to_float(aba_resumo[f"E{i}"].value) or 0.0) for i in (8, 9, 10))
        valor_s = sum((_to_float(aba_resumo[f"E{i}"].value) or 0.0) for i in (11, 12))

        wb.close()
        return (etp_display or (f"ETP {etp_esperada_norm}" if etp_esperada_norm else "")), etp_norm, gestor, preco, valor_r, valor_s

    except Exception as e:
        print(f"[Erro] ao processar anexo '{os.path.basename(path_arquivo)}': {e}")
        return "", etp_esperada_norm or "", None, None, None, None

def contar_sites_por_coluna(path_arquivo):
    try:
        wb = load_workbook(path_arquivo, data_only=True)
        aba = wb["Lista Equipamentos WDM"]

        col_inicio = None
        for col in range(1, aba.max_column + 1):
            valor = aba.cell(row=2, column=col).value
            if isinstance(valor, str) and remover_acentos(valor).strip().lower() == "total":
                col_inicio = col + 1
                break

        if col_inicio is None:
            wb.close()
            print("'Total' não encontrado na linha 2 da aba 'Lista Equipamentos WDM'.")
            return None

        colunas_com_dados = 0
        for col in range(col_inicio, aba.max_column + 1):
            for row in range(2, aba.max_row + 1):
                v = aba.cell(row=row, column=col).value
                if v is not None and str(v).strip() != "":
                    colunas_com_dados += 1
                    break

        wb.close()
        print(f"Sites detectados: {colunas_com_dados}")
        return colunas_com_dados

    except Exception as e:
        print(f"[Erro] ao contar sites no anexo '{os.path.basename(path_arquivo)}': {e}")
        return None

def coletar_estados_por_coluna(path_arquivo):
    try:
        wb = load_workbook(path_arquivo, data_only=True)
        aba = wb["Lista Equipamentos WDM"]

        col_inicio = None
        for col in range(1, aba.max_column + 1):
            valor = aba.cell(row=2, column=col).value
            if isinstance(valor, str) and remover_acentos(valor).strip().lower() == "total":
                col_inicio = col + 1
                break

        if col_inicio is None:
            wb.close()
            print("'Total' não encontrado na linha 2 da aba 'Lista Equipamentos WDM'.")
            return None

        estados_br = {
            "AC", "AL", "AP", "AM", "BA", "CE", "DF", "ES", "GO", "MA", "MT",
            "MS", "MG", "PA", "PB", "PR", "PE", "PI", "RJ", "RN", "RS", "RO",
            "RR", "SC", "SP", "SE", "TO"
        }

        estados = set()
        for col in range(col_inicio, aba.max_column + 1):
            for row in range(2, aba.max_row + 1):
                v = aba.cell(row=row, column=col).value
                if v is not None and str(v).strip() != "":
                    estado = aba.cell(row=4, column=col).value
                    if isinstance(estado, str):
                        uf = estado.strip().upper()
                        if uf in estados_br:
                            estados.add(uf)
                    break

        wb.close()
        resultado = ", ".join(sorted(estados)) if estados else None
        print(f"Estados detectados: {resultado or '(nenhum)'}")
        return resultado

    except Exception as e:
        print(f"[Erro] ao coletar estados no anexo '{os.path.basename(path_arquivo)}': {e}")
        return None

def _fmt_brl(v):
    if v is None: return "(mantido)"
    try:
        return f"R$ {float(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    except Exception:
        return str(v)

def _fmt_int(v):
    if v is None: return "(mantido)"
    try:
        vi = int(v)
        return str(vi)
    except Exception:
        return "(mantido)"

def _log_resumo(acao, etp_display, status_txt, gestor, qnt_sites, preco, valor_r, valor_s, estados, extra=""):
    print(
        f"{acao}: {etp_display} | "
        f"Status: {status_txt} | "
        f"Gestor: {gestor or '(mantido)'} | "
        f"Sites: {_fmt_int(qnt_sites)} | "
        f"Preço(Q): {_fmt_brl(preco)} | "
        f"HW+ST+MT(R): {_fmt_brl(valor_r)} | "
        f"Serviços(S): {_fmt_brl(valor_s)} | "
        f"UF: {estados or '(mantidos)'}"
        + (f" | {extra}" if extra else "")
    )


def _formatar_controle_xlsxwriter(nome_aba, linha, coluna, valor):
    if nome_aba != "Controle":
        return {}

    opcoes = {}
    if coluna < COL_LAST_BORDERED:
        opcoes["border"] = 1

    if linha > 0:
        cores = {
            0: "#FFFFFF", 1: "#FEC003", 2: "#FEC003", 3: "#FFFFFF",
            4: "#FFFFFF", 5: "#FFFFFF", 6: "#FFFFFF", 7: "#FFFFFF",
            8: "#FFFFFF", 9: "#FFFFFF", 10: "#FFFFFF", 11: "#FFFFFF",
            12: "#BFBFBF", 13: "#BFBFBF", 14: "#BFBFBF", 15: "#D6DCE4",
            16: "#BFBFBF", 17: "#FCE4D6",
        }
        if coluna in cores:
            opcoes["bg_color"] = cores[coluna]

    if coluna in (COL_PRECO - 1, COL_VAL_R - 1, COL_VAL_S - 1) and linha > 0:
        opcoes["num_format"] = BRL_ACCOUNTING
    return opcoes

# ==================== ATUALIZAÇÃO DA PLANILHA CONTROLE ====================

def atualizar_planilha(
    etp_display,
    etp_norm,
    gestor,
    status,
    preco,
    qnt_sites,
    valor_r,
    valor_s,
    estados,
    status_only: bool = False,
):
    try:
        planilhas = ler_planilhas(CAMINHO_PLANILHA_CONTROLE)
        controle = next((item for item in planilhas if item["nome"] == "Controle"), None)
        if controle is None:
            raise ValueError("A aba 'Controle' não foi encontrada na planilha.")
        linhas = controle["linhas"]
        max_row = len(linhas)
        valor_etp_col7 = etp_display if WRITE_DISPLAY_TO_SHEET else etp_norm

        def garantir_colunas(linha):
            linha.extend([None] * (COL_VAL_S - len(linha)))

        for linha in linhas:
            garantir_colunas(linha)

        # ===== Procura existente
        for row in range(1, max_row):
            linha = linhas[row]
            celula_bruta = linha[COL_ETP - 1]
            celula_etp_norm = normalizar_etp(str(celula_bruta))
            if celula_etp_norm == etp_norm:
                # merge/anti-retrocesso
                status_atual = linha[COL_STATUS - 1]
                status_final, mudou_status, motivo_status = _merge_status(status_atual, status)

                # gestor + ETP (display/norm)
                if gestor:
                    linha[COL_GESTOR - 1] = gestor
                linha[COL_ETP - 1] = valor_etp_col7

                # status
                if mudou_status:
                    linha[COL_STATUS - 1] = status_final
                    print(f"[STATUS] {motivo_status}")
                else:
                    print(f"[STATUS] Mantido: {status_atual} ({motivo_status})")

                # valores (somente se NÃO for status_only)
                if not status_only:
                    if preco is not None:
                        linha[COL_PRECO - 1] = preco
                    if isinstance(qnt_sites, (int, float)) and qnt_sites:
                        linha[COL_QNT - 1] = int(qnt_sites)
                    if valor_r is not None:
                        linha[COL_VAL_R - 1] = valor_r
                    if valor_s is not None:
                        linha[COL_VAL_S - 1] = valor_s
                    if estados:
                        linha[COL_ESTADOS - 1] = estados

                status_para_formatar = status_final if mudou_status else status_atual
                escrever_planilhas(
                    CAMINHO_PLANILHA_CONTROLE,
                    planilhas,
                    _formatar_controle_xlsxwriter,
                )

                #======== LOG ========#
                _log_resumo(
                    "Atualizado",
                    valor_etp_col7,
                    status_para_formatar or "(vazio)",
                    gestor,
                    qnt_sites if not status_only else None,
                    preco if not status_only else None,
                    valor_r if not status_only else None,
                    valor_s if not status_only else None,
                    estados if not status_only else None,
                    extra=f"status_only={status_only}"""
                )
                return "updated", row

        #======== Nova linha ========#
        nova_linha = max_row + 1
        nova_linha_dados = [None] * COL_VAL_S
        nova_linha_dados[COL_GESTOR - 1] = gestor or ""
        nova_linha_dados[COL_ETP - 1] = valor_etp_col7

        status_inserido = status or ""
        if status_inserido:
            nova_linha_dados[COL_STATUS - 1] = status_inserido

        if not status_only:
            if preco is not None:
                nova_linha_dados[COL_PRECO - 1] = preco
            if isinstance(qnt_sites, (int, float)) and qnt_sites:
                nova_linha_dados[COL_QNT - 1] = int(qnt_sites)
            if valor_r is not None:
                nova_linha_dados[COL_VAL_R - 1] = valor_r
            if valor_s is not None:
                nova_linha_dados[COL_VAL_S - 1] = valor_s
            if estados:
                nova_linha_dados[COL_ESTADOS - 1] = estados

        linhas.append(nova_linha_dados)
        escrever_planilhas(
            CAMINHO_PLANILHA_CONTROLE,
            planilhas,
            _formatar_controle_xlsxwriter,
        )

        _log_resumo(
            "Adicionado",
            valor_etp_col7,
            status_inserido or "(vazio)",
            gestor,
            qnt_sites if not status_only else None,
            preco if not status_only else None,
            valor_r if not status_only else None,
            valor_s if not status_only else None,
            estados if not status_only else None,
            extra=f"status_only={status_only}"
        )
        return "added", nova_linha

    except Exception as e:
        print(f"[Erro] ao atualizar a planilha: {e}")
        return "error", None

# ==================== LIMPAR / NORMALIZAR ====================

def limpar_e_normalizar_planilha():
    try:
        planilhas = ler_planilhas(CAMINHO_PLANILHA_CONTROLE)
        controle = next((item for item in planilhas if item["nome"] == "Controle"), None)
        if controle is None:
            raise ValueError("A aba 'Controle' não foi encontrada na planilha.")

        linhas = controle["linhas"]
        dados = linhas[1:]
        etps_vistas = {}
        for linha in dados:
            etp_norm = normalizar_etp(str(linha[COL_ETP-1]) if linha[COL_ETP-1] else "")
            versao = 1
            match = re.search(r"v(\d+)", str(linha[COL_ETP-1]).lower()) if linha[COL_ETP-1] else None
            if match:
                versao = int(match.group(1))
            if etp_norm in etps_vistas:
                if versao > etps_vistas[etp_norm][0]:
                    etps_vistas[etp_norm] = (versao, linha)
            else:
                etps_vistas[etp_norm] = (versao, linha)

        controle["linhas"] = [linhas[0]] + [linha for _, linha in etps_vistas.values()]
        escrever_planilhas(
            CAMINHO_PLANILHA_CONTROLE,
            planilhas,
            _formatar_controle_xlsxwriter,
        )
        print("Planilha normalizada e duplicadas removidas.\n")

    except Exception as e:
        print(f"[Erro] ao limpar e normalizar a planilha: {e}")

def salvar_log_execucao(adicionadas, atualizadas, mes, ano, caminho_base=None):
    try:
        timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        if not caminho_base: caminho_base = os.path.dirname(__file__)
        pasta_logs = os.path.join(caminho_base, "logs")
        os.makedirs(pasta_logs, exist_ok=True)
        path_log = os.path.join(pasta_logs, f"log_execucao_{ano}-{mes:02d}_{timestamp}.txt")

        adicionadas_unicas = sorted(set(adicionadas))
        atualizadas_unicas = sorted(set(atualizadas))

        with open(path_log, "w", encoding="utf-8") as f:
            f.write(f"Execução: {timestamp}\n")
            if DIA_FILTRADO is not None:
                f.write(f"Filtro (dia): {DIA_FILTRADO:02d}/{mes:02d}/{ano}\n")
            else:
                f.write(f"Filtro (mês): {mes:02d}/{ano}\n")
            f.write(f"Usando Outlook Restrict? {'Sim' if USE_RESTRICT else 'Não'}\n")
            f.write(f"Gravar coluna ETP como na cotação? {'Sim' if WRITE_DISPLAY_TO_SHEET else 'Não'}\n\n")

            f.write("ETPs adicionadas:\n")
            if adicionadas_unicas:
                for etp_display, linha in adicionadas_unicas:
                    f.write(f" - {etp_display} (linha {linha})\n")
            else:
                f.write(" (nenhuma)\n")

            f.write("\nETPs atualizadas:\n")
            if atualizadas_unicas:
                for etp_display, linha in atualizadas_unicas:
                    f.write(f" - {etp_display} (linha {linha})\n")
            else:
                f.write(" (nenhuma)\n")

        print(f"\nLog salvo em: {path_log}\n")
    except Exception as e:
        print(f"[Aviso] Falha ao salvar log de execução: {e}")

# ==================== FILTROS (MÊS / DIA) ====================

def _datas_restrict(mes:int, ano:int):
    inicio_dt = datetime(ano, mes, 1, 0, 0, 0)
    if mes == 12:
        fim_dt = datetime(ano + 1, 1, 1, 0, 0, 0)
    else:
        fim_dt = datetime(ano, mes + 1, 1, 0, 0, 0)
    return inicio_dt.strftime("%Y-%m-%d %H:%M"), fim_dt.strftime("%Y-%m-%d %H:%M")

def _datas_restrict_dia(dia:int, mes:int, ano:int):
    inicio_dt = datetime(ano, mes, dia, 0, 0, 0)
    fim_dt = inicio_dt + timedelta(days=1)
    return inicio_dt.strftime("%Y-%m-%d %H:%M"), fim_dt.strftime("%Y-%m-%d %H:%M")

def _iterar_emails_filtrados(inbox, mes, ano, usar_restrict=True):
    if usar_restrict:
        inicio_str, fim_str = _datas_restrict(mes, ano)
        filtro = f"[ReceivedTime] >= '{inicio_str}' AND [ReceivedTime] < '{fim_str}'"
        mensagens = inbox.Items.Restrict(filtro)
        mensagens.Sort("[ReceivedTime]", True)
        try:
            total = mensagens.Count
        except Exception:
            total = 0
        return mensagens, total
    else:
        mensagens = inbox.Items
        mensagens.Sort("[ReceivedTime]", True)
        filtradas = []
        total = mensagens.Count
        for i in range(total):
            try:
                msg = mensagens.Item(i + 1)
                dt = msg.ReceivedTime
                if dt.month == mes and dt.year == ano:
                    filtradas.append(msg)
            except Exception as e:
                print(f"[Erro] ao acessar e-mail #{i+1}: {e}")
        class _Wrapper:
            def __init__(self, data): self._data = data
            def Item(self, idx): return self._data[idx-1]
            @property
            def Count(self): return len(self._data)
        return _Wrapper(filtradas), len(filtradas)

def _iterar_emails_filtrados_dia(inbox, dia, mes, ano, usar_restrict=True):
    if usar_restrict:
        inicio_str, fim_str = _datas_restrict_dia(dia, mes, ano)
        filtro = f"[ReceivedTime] >= '{inicio_str}' AND [ReceivedTime] < '{fim_str}'"
        mensagens = inbox.Items.Restrict(filtro)
        mensagens.Sort("[ReceivedTime]", True)
        try:
            total = mensagens.Count
        except Exception:
            total = 0
        return mensagens, total
    else:
        mensagens = inbox.Items
        mensagens.Sort("[ReceivedTime]", True)
        filtradas = []
        total = mensagens.Count
        for i in range(total):
            try:
                msg = mensagens.Item(i + 1)
                dt = msg.ReceivedTime
                if dt.year == ano and dt.month == mes and dt.day == dia:
                    filtradas.append(msg)
            except Exception as e:
                print(f"[Erro] ao acessar e-mail #{i+1}: {e}")
        class _Wrapper:
            def __init__(self, data): self._data = data
            def Item(self, idx): return self._data[idx-1]
            @property
            def Count(self): return len(self._data)
        return _Wrapper(filtradas), len(filtradas)

# ==================== FUNÇÃO PRINCIPAL ====================

def processar_emails():
    global TOTAL_EMAILS_ENCONTRADOS, EMAILS_FILTRADOS, PROCESSADOS, VALIDADOS, ADICIONADAS, ATUALIZADAS
    global ADICIONADAS_LIST, ATUALIZADAS_LIST

    TOTAL_EMAILS_ENCONTRADOS = EMAILS_FILTRADOS = PROCESSADOS = VALIDADOS = ADICIONADAS = ATUALIZADAS = 0
    ADICIONADAS_LIST, ATUALIZADAS_LIST = [], []

    limpar_e_normalizar_planilha()

    inbox = testar_conexao_outlook()
    if not inbox:
        print(" Encerrando o script por erro na conexão com Outlook.")
        return

    # total da caixa (informativo)
    try:
        TOTAL_EMAILS_ENCONTRADOS = inbox.Items.Count
    except Exception:
        TOTAL_EMAILS_ENCONTRADOS = 0

    # Seleção do modo de filtro
    if DIA_FILTRADO is not None:
        mensagens, EMAILS_FILTRADOS = _iterar_emails_filtrados_dia(inbox, DIA_FILTRADO, MES_FILTRADO, ANO_FILTRADO, USE_RESTRICT)
        print(f"\n Total de e-mails (Inbox): {TOTAL_EMAILS_ENCONTRADOS}")
        print(f" E-mails filtrados para {DIA_FILTRADO:02d}/{MES_FILTRADO:02d}/{ANO_FILTRADO}: {EMAILS_FILTRADOS}")
    else:
        mensagens, EMAILS_FILTRADOS = _iterar_emails_filtrados(inbox, MES_FILTRADO, ANO_FILTRADO, USE_RESTRICT)
        print(f"\n Total de e-mails (Inbox): {TOTAL_EMAILS_ENCONTRADOS}")
        print(f" E-mails filtrados para {MES_FILTRADO:02d}/{ANO_FILTRADO}: {EMAILS_FILTRADOS}")

    print(f" Usando Outlook Restrict? {'Sim' if USE_RESTRICT else 'Não'}")
    print(" Iniciando processamento dos e-mails filtrados...\n")

    bloco_tamanho = 200

    for i in range(EMAILS_FILTRADOS):
        anexos = []
        try:
            msg = mensagens.Item(i + 1)
            assunto = msg.Subject or ""
            corpo = msg.Body or ""
            print(f"Analisando e-mail: {assunto}")

            anexos = salvar_anexos(msg)

            # Detecta T2 e qual é o anexo T2 (se houver)
            tem_t2, etp_norm_t2, path_t2 = detectar_anexo_t2(anexos)

            if not validar_email(assunto, anexos, corpo):
                print("E-mail não se enquadra nos critérios (assunto, anexos ou chave). Pulando...\n")
                PROCESSADOS += 1
                continue

            gestor = None
            preco = None
            qnt_sites = None
            valor_r = None
            valor_s = None
            estados = None

            VALIDADOS += 1
            PROCESSADOS += 1

            # Determina ETP
            etp_norm = extrair_etp_do_assunto(assunto)
            if not etp_norm and tem_t2:
                etp_norm = etp_norm_t2
            etp_display = f"ETP {etp_norm}" if etp_norm else ""

            # Vamos decidir se há anexo de COTAÇÃO (para ler dados). Se for só T2, NÃO lemos nada.
            def is_cotacao(nome_base: str) -> bool:
                nb = remover_acentos(nome_base.lower())
                return ("cotacao" in nb) or (("etp" in nb) and not nb.startswith("t2_etp"))

            anexo_cotacao = None
            if anexos:
                candidatos = [p for p in anexos if is_cotacao(os.path.basename(p))]
                if candidatos:
                    def score(p):
                        nome = remover_acentos(os.path.basename(p).lower())
                        v = re.search(r"\bv(\d+)\b", nome)
                        ver = int(v.group(1)) if v else 1
                        bonus = 0
                        if "cotacao" in nome and "etp" in nome:
                            bonus += 1000
                        return bonus + ver
                    anexo_cotacao = sorted(candidatos, key=score, reverse=True)[0]

            # Se tem anexo de cotação, extrai dados
            if anexo_cotacao:
                etp_disp_tmp, etp_norm_tmp, gestor_anexo, preco_anexo, vr_tmp, vs_tmp = extrair_dados_anexo(anexo_cotacao, etp_norm)
                etp_display = etp_disp_tmp or etp_display
                etp_norm = etp_norm_tmp or etp_norm
                gestor = gestor_anexo if gestor_anexo is not None else gestor
                preco = preco_anexo if preco_anexo is not None else preco
                valor_r = vr_tmp if vr_tmp is not None else valor_r
                valor_s = vs_tmp if vs_tmp is not None else valor_s
                qnt_sites = contar_sites_por_coluna(anexo_cotacao) if qnt_sites is None else qnt_sites
                estados = coletar_estados_por_coluna(anexo_cotacao) if estados is None else estados

            if not etp_norm:
                print("ETP não encontrada nesse e-mail. Pulando...\n")
                continue

            # Define status
            status = identificar_status_etp(assunto, anexos, corpo, msg.ReceivedTime, tem_t2=tem_t2)
            if not status:
                print(f"[Aviso] Status não identificado para ETP {etp_display or etp_norm}. Pulando...\n")
                continue

            # Se é só T2 (sem cotação), atualiza SOMENTE o status
            status_only = (tem_t2 and anexo_cotacao is None)

            result, linha = atualizar_planilha(
                etp_display or etp_norm,
                etp_norm,
                gestor,
                status,
                preco,
                qnt_sites,
                valor_r,
                valor_s,
                estados,
                status_only=status_only
            )

            # >>> AQUI ENTRA O BOT QUANTITATIVO (somente se tiver cotação) <<<
            if result in ("updated", "added") and (not status_only) and anexo_cotacao:
                try:
                    print(f"[Quant] Chamando Bot_Quantitativo para {etp_display or etp_norm}...")
                    processar_quantitativo(
                        CAMINHO_PLANILHA_CONTROLE,
                        etp_display or etp_norm,
                        etp_norm,
                        gestor,
                        anexo_cotacao,
                        base_dir_itens=path,  # <<< garante que ele ache o itens.txt
                    )
                    print(f"[Quant] Finalizado para {etp_display or etp_norm}.")
                except Exception as e:
                    print(f"[Quant][ERRO] Falha ao rodar quantitativo: {e}")


            if result == "updated":
                ATUALIZADAS += 1
                ATUALIZADAS_LIST.append((etp_display or f"ETP {etp_norm}", linha))
            elif result == "added":
                ADICIONADAS += 1
                ADICIONADAS_LIST.append((etp_display or f"ETP {etp_norm}", linha))

            # marcar como lido
            try:
                msg.Unread = False
            except Exception:
                pass

            if PROCESSADOS % bloco_tamanho == 0:
                print(f"Bloco de {bloco_tamanho} e-mails processado.")
                print(f"Total analisados até agora: {PROCESSADOS}")
                print(f"Total validados (com ETP ou anexos válidos): {VALIDADOS}\n")
                print("Pausando 5 segundos...\n")
                time.sleep(5)

        except Exception as e:
            print(f"[Erro] ao processar e-mail: {e}")
        finally:
            for p in anexos:
                try:
                    os.remove(p)
                except:
                    pass

    print("Processamento concluído!")
    print(f"Total de e-mails analisados: {PROCESSADOS}")
    print(f"Total de e-mails validados (processados com dados): {VALIDADOS}\n")

    # Resumo no console
    adicionadas_unicas = sorted(set(ADICIONADAS_LIST))
    atualizadas_unicas = sorted(set(ATUALIZADAS_LIST))
    print("ETPs ADICIONADAS:")
    for etp_display, linha in adicionadas_unicas:
        print(f" - {etp_display} (linha {linha})")
    print("\nETPs ATUALIZADAS:")
    for etp_display, linha in atualizadas_unicas:
        print(f" - {etp_display} (linha {linha})")

    salvar_log_execucao(ADICIONADAS_LIST, ATUALIZADAS_LIST, MES_FILTRADO, ANO_FILTRADO)

# ==================== INTERFACE ====================

def iniciar_interface():
    class BotEmailApp:
        def __init__(self, root):
            self.root = root
            self.root.title("Automação para Coleta de E-mails")
            try:
                if os.path.exists(path_icone) and path_icone.lower().endswith('.ico'):
                    self.root.iconbitmap(path_icone)
            except Exception:
                pass
            self.create_widgets()

        def create_widgets(self):
            rowi = 0
            frame = ttk.Frame(self.root, padding="10")
            frame.grid(row=0, column=0, sticky="nw")

            ttk.Label(frame, text="Mês (ex: 6 para Junho):").grid(row=rowi, column=0, sticky="w")
            self.mes_entry = ttk.Entry(frame, width=8)
            self.mes_entry.grid(row=rowi, column=1, sticky="w", padx=(6,0))
            self.mes_entry.insert(0, "6")
            rowi += 1

            ttk.Label(frame, text="Ano (ex: 2025):").grid(row=rowi, column=0, sticky="w")
            self.ano_entry = ttk.Entry(frame, width=8)
            self.ano_entry.grid(row=rowi, column=1, sticky="w", padx=(6,0))
            self.ano_entry.insert(0, "2025")
            rowi += 1

            ttk.Label(frame, text="Dia (ex: 1-31):").grid(row=rowi, column=0, sticky="w")
            self.dia_entry = ttk.Entry(frame, width=8)
            self.dia_entry.grid(row=rowi, column=1, sticky="w", padx=(6,0))
            # em branco = filtro mensal
            rowi += 1

            # Checkbox: grava coluna ETP como está na cotação (display) ou normalizada
            self.var_write_display = tk.BooleanVar(value=True)
            self.chk_display = ttk.Checkbutton(
                frame,
                text="Gravar coluna ETP exatamente como na cotação (ex.: ETP 1234-5678 V3)",
                variable=self.var_write_display
            )
            self.chk_display.grid(row=rowi, column=0, columnspan=2, sticky="w", pady=(6, 0))
            rowi += 1

            # Checkbox: usar Outlook Restrict (rápido) ou varrer tudo (compatibilidade)
            self.var_use_restrict = tk.BooleanVar(value=True)
            self.chk_restrict = ttk.Checkbutton(
                frame,
                text="Usar filtro Outlook Restrict (mais rápido)",
                variable=self.var_use_restrict
            )
            self.chk_restrict.grid(row=rowi, column=0, columnspan=2, sticky="w", pady=(2, 8))
            rowi += 1

            self.start_button = ttk.Button(frame, text="Iniciar Processamento", command=self.iniciar_processamento)
            self.start_button.grid(row=rowi, column=0, columnspan=2, pady=8)
            rowi += 1

            self.resultados = tk.Text(frame, height=22, width=100, state="disabled")
            self.resultados.grid(row=rowi, column=0, columnspan=2, pady=(6,0))
            rowi += 1

            self.rodape = ttk.Label(
                frame,
                text="Desenvolvido por Gustavo Vieira (g50046141)\nManutenção Antonio Satiro (a50056403)",
                font=("Arial", 8),
                foreground="gray"
            )
            self.rodape.grid(row=rowi, column=0, columnspan=2, pady=(10, 0), sticky="e")

        def iniciar_processamento(self):
            try:
                global MES_FILTRADO, ANO_FILTRADO, DIA_FILTRADO, WRITE_DISPLAY_TO_SHEET, USE_RESTRICT
                MES_FILTRADO = int(self.mes_entry.get())
                ANO_FILTRADO = int(self.ano_entry.get())

                # dia opcional
                dia_raw = (self.dia_entry.get() or "").strip()
                DIA_FILTRADO = int(dia_raw) if dia_raw.isdigit() else None

                WRITE_DISPLAY_TO_SHEET = bool(self.var_write_display.get())
                USE_RESTRICT = bool(self.var_use_restrict.get())

                self.resultados.config(state="normal")
                self.resultados.delete(1.0, tk.END)
                self.resultados.insert(tk.END, "Iniciando processamento...\n")
                if DIA_FILTRADO is not None:
                    self.resultados.insert(tk.END, f"Filtro diário: {DIA_FILTRADO:02d}/{MES_FILTRADO:02d}/{ANO_FILTRADO}\n")
                else:
                    self.resultados.insert(tk.END, f"Filtro mensal: {MES_FILTRADO:02d}/{ANO_FILTRADO}\n")
                self.resultados.insert(tk.END, f"Gravar coluna ETP como na cotação: {'Sim' if WRITE_DISPLAY_TO_SHEET else 'Não'}\n")
                self.resultados.insert(tk.END, f"Usar Outlook Restrict: {'Sim' if USE_RESTRICT else 'Não'}\n\n")
                self.resultados.config(state="disabled")

                thread = threading.Thread(target=self.executar_bot, daemon=True)
                thread.start()

            except Exception as e:
                messagebox.showerror("Erro", str(e))

        def executar_bot(self):
            processar_emails()
            self.root.after(0, self._mostrar_resumo)

        def _mostrar_resumo(self):
            adicionadas_unicas = sorted(set(ADICIONADAS_LIST))
            atualizadas_unicas = sorted(set(ATUALIZADAS_LIST))

            self.resultados.config(state="normal")
            self.resultados.insert(tk.END, "Processamento Finalizado\n")
            self.resultados.insert(tk.END, f"Total de e-mails (Inbox): {TOTAL_EMAILS_ENCONTRADOS}\n")
            if DIA_FILTRADO is not None:
                self.resultados.insert(tk.END, f"E-mails filtrados (dia): {DIA_FILTRADO:02d}/{MES_FILTRADO:02d}/{ANO_FILTRADO}\n")
            else:
                self.resultados.insert(tk.END, f"E-mails filtrados (mês): {MES_FILTRADO:02d}/{ANO_FILTRADO}\n")
            self.resultados.insert(tk.END, f"E-mails processados: {PROCESSADOS}\n")
            self.resultados.insert(tk.END, f"E-mails validados: {VALIDADOS}\n")
            self.resultados.insert(tk.END, f"ETPs adicionadas: {ADICIONADAS}\n")
            self.resultados.insert(tk.END, f"ETPs atualizadas: {ATUALIZADAS}\n")

            self.resultados.insert(tk.END, "\nETPs adicionadas:\n")
            if adicionadas_unicas:
                for etp_display, linha in adicionadas_unicas:
                    self.resultados.insert(tk.END, f"- {etp_display} (linha {linha})\n")
            else:
                self.resultados.insert(tk.END, "(nenhuma)\n")

            self.resultados.insert(tk.END, "\nETPs atualizadas:\n")
            if atualizadas_unicas:
                for etp_display, linha in atualizadas_unicas:
                    self.resultados.insert(tk.END, f"- {etp_display} (linha {linha})\n")
            else:
                self.resultados.insert(tk.END, "(nenhuma)\n")

            self.resultados.insert(tk.END, "\nUm arquivo de log foi salvo na pasta 'logs' ao lado do script.\n")
            self.resultados.config(state="disabled")

    root = tk.Tk()
    app = BotEmailApp(root)
    root.mainloop()

if __name__ == "__main__":
    iniciar_interface()