from __future__ import annotations

import os
from typing import Callable, Any

import xlsxwriter
from openpyxl import load_workbook


FormatCallback = Callable[[str, int, int, Any], dict]


def ler_planilhas(caminho: str) -> list[dict]:
    """Lê os valores das abas sem usar openpyxl para gravar o arquivo."""
    wb = load_workbook(caminho, data_only=False)
    try:
        planilhas = []
        for ws in wb.worksheets:
            linhas = [
                [cell.value for cell in row]
                for row in ws.iter_rows(
                    min_row=1,
                    max_row=max(ws.max_row, 1),
                    min_col=1,
                    max_col=max(ws.max_column, 1),
                )
            ]
            planilhas.append({"nome": ws.title, "linhas": linhas})
        return planilhas
    finally:
        wb.close()


def escrever_planilhas(
    caminho: str,
    planilhas: list[dict],
    format_callback: FormatCallback | None = None,
) -> None:
    """Cria o xlsx com xlsxwriter a partir dos valores lidos/ajustados."""
    temporario = f"{caminho}.xlsxwriter.tmp"
    workbook = xlsxwriter.Workbook(temporario)
    formatos = {}

    def obter_formato(opcoes: dict):
        chave = tuple(sorted(opcoes.items()))
        if chave not in formatos:
            formatos[chave] = workbook.add_format(opcoes)
        return formatos[chave]

    try:
        for dados in planilhas:
            worksheet = workbook.add_worksheet(dados["nome"][:31])
            for row_idx, linha in enumerate(dados["linhas"]):
                for col_idx, valor in enumerate(linha):
                    opcoes = format_callback(dados["nome"], row_idx, col_idx, valor) if format_callback else {}
                    formato = obter_formato(opcoes) if opcoes else None
                    if isinstance(valor, str) and valor.startswith("="):
                        worksheet.write_formula(row_idx, col_idx, valor, formato)
                    elif valor is None:
                        worksheet.write_blank(row_idx, col_idx, None, formato)
                    else:
                        worksheet.write(row_idx, col_idx, valor, formato)
        workbook.close()
        os.replace(temporario, caminho)
    except Exception:
        workbook.close()
        if os.path.exists(temporario):
            os.remove(temporario)
        raise